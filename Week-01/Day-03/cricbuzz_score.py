from playwright.sync_api import sync_playwright, TimeoutError as PlaywrightTimeoutError
from openpyxl import Workbook
from datetime import datetime
import os


# -----------------------------------
# Configuration
# -----------------------------------

MATCH_URL = (
    "https://www.cricbuzz.com/live-cricket-scores/"
    "171070/ind-vs-pak-gold-medal-match-asian-games-2026"
)

SCORE_SELECTOR = "#sticky-mcomplete"

EXCEL_FILE = "cricket_score.xlsx"
SCREENSHOT_FILE = "score.png"


# -----------------------------------
# Start Playwright
# -----------------------------------

with sync_playwright() as p:

    browser = p.chromium.launch(headless=False)

    page = browser.new_page()

    print("Opening Cricbuzz...")

    page.goto(
        MATCH_URL,
        wait_until="domcontentloaded"
    )

    try:

        print("Waiting for score...")

        # Wait for the actual score element
        page.wait_for_selector(
            SCORE_SELECTOR,
            state="visible",
            timeout=30000
        )

        # Read score text
        score = page.locator(
            SCORE_SELECTOR
        ).inner_text()

        print("\nLive Cricket Score:")
        print("-------------------")
        print(score)

        # -----------------------------------
        # Screenshot
        # -----------------------------------

        page.screenshot(
            path=SCREENSHOT_FILE,
            full_page=True
        )

        print("\nScreenshot saved as:", SCREENSHOT_FILE)

        # -----------------------------------
        # Prepare Excel
        # -----------------------------------

        wb = Workbook()

        ws = wb.active
        ws.title = "Cricket Score"

        # Headings
        ws["A1"] = "Date"
        ws["B1"] = "Time"
        ws["C1"] = "Match"
        ws["D1"] = "Score"

        # Current date and time
        now = datetime.now()

        ws["A2"] = now.strftime("%d-%m-%Y")
        ws["B2"] = now.strftime("%H:%M:%S")

        ws["C2"] = "PAK vs IND - Gold Medal Match"

        ws["D2"] = score

        # Improve column width
        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 35
        ws.column_dimensions["D"].width = 50

        # Wrap multiline score
        ws["D2"].alignment = ws["D2"].alignment.copy(
            wrap_text=True
        )

        # Save Excel
        wb.save(EXCEL_FILE)

        print("Excel file saved as:", EXCEL_FILE)

        print("\nFiles saved in:")
        print(os.getcwd())

    except PlaywrightTimeoutError:

        print(
            "Score element was not found within 30 seconds."
        )

        print(
            "Check whether the match page or selector has changed."
        )

    finally:

        browser.close()